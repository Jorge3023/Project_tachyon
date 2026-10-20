"""Exportación de datos a CSV y Excel: filas filtradas del archivo y/o los números detrás de cada gráfica.

Las imágenes, el PDF y el PowerPoint se generan en el navegador (lo que se exporta es lo que se ve); aquí solo
se arma lo que depende de los datos completos, que el navegador nunca recibe.
"""
import io
import re
from datetime import datetime, timezone
from typing import Optional

import pandas as pd

MAX_FILAS_CSV = 500_000
MAX_FILAS_XLSX = 50_000           # openpyxl escribe celda por celda: más filas tardan y consumen memoria en un hosting pequeño
_RX_PELIGRO = re.compile(r"^[=+\-@\t\r]")
_RX_NUMERO_TEXTO = re.compile(r"^[+-]?\d+([.,]\d+)?$")
_RX_ILEGAL_HOJA = re.compile(r"[\[\]:*?/\\]")
_RX_ILEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _texto_seguro(s: pd.Series) -> pd.Series:
    """Evita la inyección de fórmulas: un texto que empieza con = + - @ se abriría como fórmula en Excel."""
    s = s.astype("string")
    peligro = s.str.match(_RX_PELIGRO, na=False) & ~s.str.match(_RX_NUMERO_TEXTO, na=False)
    return s.mask(peligro, "'" + s)


def preparar(df: pd.DataFrame, esquema: list) -> pd.DataFrame:
    """Columnas del esquema (sin las ignoradas) en su orden, con textos seguros y fechas como fecha."""
    cols = [c["nombre"] for c in esquema if c.get("rol") != "ignorar" and c["nombre"] in df.columns]
    out = df[cols].copy()
    roles = {c["nombre"]: c["rol"] for c in esquema}
    for c in cols:
        if roles.get(c) == "fecha" and pd.api.types.is_datetime64_any_dtype(out[c]):
            continue
        if roles.get(c) != "numero" and not pd.api.types.is_datetime64_any_dtype(out[c]):
            out[c] = _texto_seguro(out[c])
    return out


def descripcion_filtros(config: dict) -> list:
    """Filtros globales de una configuración, en texto legible."""
    out = []
    for f in config.get("filtros", []):
        if "valores" in f:
            vals = ", ".join(f["valores"][:6]) + ("…" if len(f["valores"]) > 6 else "")
            out.append(f"{f['col']}: {vals}")
        elif "desde" in f:
            out.append(f"{f['col']}: {f.get('desde') or '…'} a {f.get('hasta') or '…'}")
        else:
            lo, hi = f.get("min"), f.get("max")
            out.append(f"{f['col']}: {'' if lo is None else lo} a {'' if hi is None else hi}")
    return out


# ── CSV ─────────────────────────────────────────────────────────────────────────────────────────────
def a_csv(df: pd.DataFrame) -> bytes:
    """UTF-8 con BOM (Excel en español lo abre con los acentos bien), coma como separador."""
    d = df.copy()
    for c in d.columns:
        if pd.api.types.is_datetime64_any_dtype(d[c]):
            solo_fecha = (d[c].dropna().dt.normalize() == d[c].dropna()).all()
            d[c] = d[c].dt.strftime("%Y-%m-%d" if solo_fecha else "%Y-%m-%d %H:%M:%S")
    return ("\ufeff" + d.to_csv(index=False, lineterminator="\r\n")).encode("utf-8")


# ── Tablas con los números de cada gráfica ─────────────────────────────────────────────────────────
def tabla_widget(w: dict, r: dict):
    """(columnas, filas) con los números que dibuja una gráfica; None si no aplica (o si hubo error)."""
    if not r or r.get("error"):
        return None
    t = w["tipo"]
    if t == "kpi":
        filas = [["Valor", r.get("valor")]]
        d = r.get("delta")
        if d:
            filas += [["Variación contra el periodo anterior (%)", d.get("pct")], ["Periodo actual", d.get("actual")],
                      ["Periodo anterior", d.get("anterior")]]
        return ["Indicador", w.get("titulo") or "Valor"], filas
    if t in ("linea", "area", "barras"):
        ser = r.get("series", [])
        return [w.get("x") or "Categoría"] + [s["nombre"] for s in ser], \
            [[c] + [s["valores"][i] for s in ser] for i, c in enumerate(r.get("categorias", []))]
    if t == "dona":
        return [w.get("x") or "Categoría", "Valor", "%"], [[p["nombre"], p["valor"], p.get("pct")] for p in r.get("partes") or []]
    if t == "pareto":
        return [w.get("x") or "Categoría", "Valor", "% acumulado"], \
            [[c, r["valores"][i], r["acumulado"][i]] for i, c in enumerate(r.get("categorias", []))]
    if t == "dispersion":
        nombres = r.get("series", [])
        cols = [w.get("x") or "X", w.get("y") or "Y"] + (["Serie"] if len(nombres) > 1 else [])
        return cols, [[p[0], p[1]] + ([nombres[p[2]]] if len(nombres) > 1 else []) for p in r.get("puntos", [])]
    if t == "histograma":
        return ["Desde", "Hasta", "Filas"], [[b["desde"], b["hasta"], b["n"]] for b in r.get("bins", [])]
    if t == "calor":
        return [w.get("y") or "Fila"] + r.get("columnas", []), [[f] + r["valores"][j] for j, f in enumerate(r.get("filas", []))]
    if t == "mapa":
        nombres = r.get("series", [])
        med = r.get("medida")
        cols = ["Latitud", "Longitud", f"{r.get('agg', '')} de {med}" if med else "Filas"] + (["Filas"] if med else []) \
            + (["Categoría"] if nombres else [])
        return cols, [[p[0], p[1], p[2] if med else p[3]] + ([p[3]] if med else []) + ([nombres[p[4]]] if nombres else [])
                      for p in r.get("puntos", [])]
    if t == "tabla":
        return list(r.get("columnas", [])), r.get("filas", [])
    return None


def _nombre_hoja(titulo: str, usados: set) -> str:
    base = _RX_ILEGAL_HOJA.sub(" ", titulo or "Gráfica").strip().strip("'") or "Gráfica"
    base = base[:28]
    n, nombre = 1, base
    while nombre.lower() in usados:
        n += 1
        nombre = f"{base[:28 - len(str(n)) - 1]} {n}"
    usados.add(nombre.lower())
    return nombre


def _limpio(v):
    """Valor apto para una celda: sin caracteres de control, con tipos básicos de Python."""
    if isinstance(v, str):
        return _RX_ILEGAL_XML.sub("", v)
    if isinstance(v, float) and v != v:
        return None
    return v


# ── Excel ───────────────────────────────────────────────────────────────────────────────────────────
def a_xlsx(df: Optional[pd.DataFrame], hojas_graficas: list, resumen: list) -> bytes:
    """
    df: filas filtradas (o None para no incluir la hoja «Datos»).
    hojas_graficas: [(título, columnas, filas)].  resumen: [(etiqueta, valor)].
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    negrita = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F4E8C")

    def encabezado(ws, cols):
        celdas = []
        for c in cols:
            x = WriteOnlyCell(ws, value=_limpio(str(c)))
            x.font, x.fill, x.alignment = negrita, fondo, Alignment(vertical="center")
            celdas.append(x)
        ws.append(celdas)

    def anchos(ws, cols, muestra):
        for i, c in enumerate(cols):
            largo = max([len(str(c))] + [len(str(f[i])) for f in muestra[:200] if i < len(f) and f[i] is not None])
            ws.column_dimensions[get_column_letter(i + 1)].width = min(max(largo + 2, 8), 48)

    ws = wb.create_sheet("Resumen")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    for etiqueta, valor in resumen:
        a = WriteOnlyCell(ws, value=_limpio(str(etiqueta)))
        a.font = Font(bold=True)
        ws.append([a, _limpio(valor)])

    usados = {"resumen"}
    if df is not None:
        ws = wb.create_sheet("Datos")
        usados.add("datos")
        cols = list(df.columns)
        filas = df.astype(object).where(df.notna(), None).itertuples(index=False, name=None)
        muestra = []
        anchos_listos = False
        ws.freeze_panes = "A2"
        encabezado(ws, cols)
        n = 0
        for fila in filas:
            if not anchos_listos:
                muestra.append(fila)
                if len(muestra) >= 200:
                    anchos(ws, cols, muestra); anchos_listos = True
            ws.append([_limpio(v) for v in fila])
            n += 1
        if not anchos_listos:
            anchos(ws, cols, muestra)
        ws.auto_filter.ref = f"A1:{get_column_letter(max(len(cols), 1))}{n + 1}"
    for titulo, cols, filas in hojas_graficas:
        ws = wb.create_sheet(_nombre_hoja(titulo, usados))
        ws.freeze_panes = "A2"
        encabezado(ws, cols)
        anchos(ws, cols, filas)
        for f in filas:
            ws.append([_limpio(v) for v in f])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def nombre_archivo(base: str, ext: str) -> str:
    """Nombre seguro para descargar: sin rutas ni caracteres raros, con la fecha."""
    import unicodedata
    base = unicodedata.normalize("NFKD", base or "datos").encode("ascii", "ignore").decode()
    base = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", base)
    base = re.sub(r"_+", "_", re.sub(r"[^A-Za-z0-9._-]+", "_", base)).strip("._-")[:60] or "datos"
    return f"{base}_{datetime.now(timezone.utc).strftime('%Y-%m-%d')}.{ext}"
